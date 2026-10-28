from __future__ import annotations

import csv
import hashlib
import io
import json
import mimetypes
import os
import re
import secrets
import sqlite3
import uuid
from datetime import date, datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from urllib.parse import quote

from flask import (
    Flask, abort, flash, g, jsonify, make_response, redirect, render_template,
    request, send_file, session, url_for,
)
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
INSTANCE_DIR = BASE_DIR / "instance"
UPLOAD_DIR = INSTANCE_DIR / "uploads"
DB_PATH = INSTANCE_DIR / "aegisai.sqlite3"
RULES_PATH = INSTANCE_DIR / "rules.json"
DEFAULT_RULES_PATH = BASE_DIR / "config" / "default_rules.json"
REQUIRED_DOCS_PATH = BASE_DIR / "config" / "required_documents.json"

INSTANCE_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("AEGISAI_SECRET_KEY") or secrets.token_hex(32)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024
app.config["MAX_FILES_PER_APPLICATION"] = 12
app.config["UPLOAD_FOLDER"] = str(UPLOAD_DIR)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.environ.get("AEGISAI_HTTPS", "0") == "1"

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "csv", "png", "jpg", "jpeg", "tif", "tiff", "webp"}
SEVERITY_WEIGHT = {"Critical": 25, "High": 14, "Medium": 7, "Low": 3, "Info": 0}
SEVERITY_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "Info": 4}
ISSUER_DOC_TYPES = {
    "Extension of Approval (EoA) letter", "Accreditation / recognition certificate",
    "NOC / permission", "Fire safety certificate", "Building stability certificate",
    "TER payment receipt", "Faculty qualification evidence",
}

SYNTHETIC_ISSUER_REGISTRY = {
    "DEMO-EOA-001": {
        "status": "Verified (synthetic source)",
        "institution_code": "AICTE-DEMO-042",
        "valid_until": "31/03/2027",
        "issuer": "AegisAI Synthetic Approval Registry",
        "note": "Reference and institution code match the synthetic test record.",
    },
    "DEMO-ACC-EXPIRED": {
        "status": "Expired (synthetic source)",
        "institution_code": "AICTE-DEMO-042",
        "valid_until": "31/03/2025",
        "issuer": "AegisAI Synthetic Accreditation Registry",
        "note": "Synthetic test record is past its validity date.",
    },
    "DEMO-TER-001": {
        "status": "Verified (synthetic source)",
        "institution_code": "AICTE-DEMO-042",
        "valid_until": "31/03/2027",
        "issuer": "AegisAI Synthetic Payment Registry",
        "note": "Synthetic payment reference exists and matches the institution code.",
    },
    "DEMO-NOC-MISMATCH": {
        "status": "Mismatch (synthetic source)",
        "institution_code": "AICTE-DEMO-999",
        "valid_until": "31/03/2027",
        "issuer": "AegisAI Synthetic NOC Registry",
        "note": "Synthetic source returned a different institution code.",
    },
}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def load_json(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


def get_required_documents() -> dict[str, list[str]]:
    return load_json(REQUIRED_DOCS_PATH, {})


def get_rules() -> dict[str, bool]:
    defaults = load_json(DEFAULT_RULES_PATH, {})
    overrides = load_json(RULES_PATH, {})
    return {key: bool(overrides.get(key, value)) for key, value in defaults.items()}


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_exc: BaseException | None = None) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db() -> None:
    conn = sqlite3.connect(DB_PATH)
    conn.executescript(
        """
        PRAGMA foreign_keys = ON;
        CREATE TABLE IF NOT EXISTS applications (
            id TEXT PRIMARY KEY,
            institution_name TEXT NOT NULL,
            institution_code TEXT NOT NULL,
            category TEXT NOT NULL,
            academic_year TEXT NOT NULL,
            contact_email TEXT DEFAULT '',
            status TEXT NOT NULL DEFAULT 'Draft',
            readiness_score INTEGER NOT NULL DEFAULT 100,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            is_demo INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            application_id TEXT NOT NULL REFERENCES applications(id) ON DELETE CASCADE,
            doc_type TEXT NOT NULL,
            original_name TEXT NOT NULL,
            stored_name TEXT,
            extension TEXT NOT NULL DEFAULT '',
            mime_type TEXT NOT NULL DEFAULT 'application/octet-stream',
            size_bytes INTEGER NOT NULL DEFAULT 0,
            sha256 TEXT NOT NULL,
            extracted_text TEXT NOT NULL DEFAULT '',
            extraction_status TEXT NOT NULL DEFAULT 'Not processed',
            metadata_json TEXT NOT NULL DEFAULT '{}',
            source_status TEXT NOT NULL DEFAULT 'Not checked',
            source_name TEXT NOT NULL DEFAULT '',
            source_detail TEXT NOT NULL DEFAULT '',
            issuer_reference TEXT NOT NULL DEFAULT '',
            signature_ok INTEGER NOT NULL DEFAULT 1,
            is_demo INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS findings (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            application_id TEXT NOT NULL REFERENCES applications(id) ON DELETE CASCADE,
            finding_key TEXT NOT NULL,
            doc_id INTEGER REFERENCES documents(id) ON DELETE SET NULL,
            severity TEXT NOT NULL,
            title TEXT NOT NULL,
            detail TEXT NOT NULL,
            evidence TEXT NOT NULL DEFAULT '',
            recommendation TEXT NOT NULL DEFAULT '',
            status TEXT NOT NULL DEFAULT 'Open',
            is_active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            UNIQUE(application_id, finding_key)
        );
        CREATE TABLE IF NOT EXISTS audit_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            application_id TEXT REFERENCES applications(id) ON DELETE CASCADE,
            actor TEXT NOT NULL DEFAULT 'Demo reviewer',
            action TEXT NOT NULL,
            detail TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_apps_updated ON applications(updated_at DESC);
        CREATE INDEX IF NOT EXISTS idx_docs_app ON documents(application_id);
        CREATE INDEX IF NOT EXISTS idx_findings_app ON findings(application_id, is_active);
        CREATE INDEX IF NOT EXISTS idx_audit_app ON audit_logs(application_id, created_at DESC);
        """
    )
    conn.commit()
    conn.close()


init_db()


def record_audit(db: sqlite3.Connection, app_id: str | None, action: str, detail: str = "", actor: str = "Demo reviewer") -> None:
    db.execute(
        "INSERT INTO audit_logs(application_id, actor, action, detail, created_at) VALUES (?, ?, ?, ?, ?)",
        (app_id, actor, action, detail[:2000], utc_now()),
    )


def csrf_token() -> str:
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(24)
    return str(session["csrf_token"])


@app.context_processor
def inject_template_helpers() -> dict[str, Any]:
    return {"csrf_token": csrf_token, "current_year": date.today().year}


@app.before_request
def protect_post_requests() -> None:
    if request.method == "POST":
        sent = request.form.get("csrf_token") or request.headers.get("X-CSRF-Token", "")
        expected = session.get("csrf_token", "")
        if not expected or not secrets.compare_digest(str(sent), str(expected)):
            abort(400, description="Your form session expired. Refresh the page and try again.")


@app.after_request
def set_security_headers(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
    response.headers.setdefault("Cache-Control", "no-store")
    return response


def normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").casefold())


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def file_signature_ok(ext: str, data: bytes) -> bool:
    ext = ext.lower()
    if ext == "pdf":
        return data.lstrip().startswith(b"%PDF-")
    if ext in {"png"}:
        return data.startswith(b"\x89PNG\r\n\x1a\n")
    if ext in {"jpg", "jpeg"}:
        return data.startswith(b"\xff\xd8\xff")
    if ext in {"tif", "tiff"}:
        return data.startswith((b"II*\x00", b"MM\x00*"))
    if ext in {"docx"}:
        return data.startswith(b"PK\x03\x04")
    if ext in {"webp"}:
        return len(data) > 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP"
    return True


def extract_document(data: bytes, ext: str) -> tuple[str, str, dict[str, Any]]:
    text = ""
    metadata: dict[str, Any] = {}
    try:
        if ext in {"txt", "csv"}:
            text = data.decode("utf-8-sig", errors="replace")
            return text[:150_000], "Text extracted", metadata
        if ext == "docx":
            from docx import Document
            document = Document(io.BytesIO(data))
            chunks = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(chunks)
            return text[:150_000], ("Text extracted" if text.strip() else "No readable text"), metadata
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data), strict=False)
            meta = reader.metadata
            if meta:
                metadata = {
                    "title": str(meta.title or "")[:300],
                    "author": str(meta.author or "")[:200],
                    "creator": str(meta.creator or "")[:200],
                    "producer": str(meta.producer or "")[:200],
                    "creation_date": str(meta.creation_date or ""),
                    "modification_date": str(meta.modification_date or ""),
                }
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:30])
            if text.strip():
                return text[:150_000], "Text extracted", metadata
            # OCR is attempted only when the optional OCR engine and PDF renderer are available.
            try:
                import fitz
                import pytesseract
                from PIL import Image
                pdf = fitz.open(stream=data, filetype="pdf")
                ocr_chunks = []
                for page in pdf[:3]:
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                    image = Image.open(io.BytesIO(pix.tobytes("png")))
                    ocr_chunks.append(pytesseract.image_to_string(image))
                ocr_text = "\n".join(ocr_chunks)
                if ocr_text.strip():
                    return ocr_text[:150_000], "OCR extracted", metadata
                return "", "No readable text", metadata
            except Exception:
                return "", "OCR unavailable for scanned PDF", metadata
        if ext in {"png", "jpg", "jpeg", "tif", "tiff", "webp"}:
            try:
                import pytesseract
                from PIL import Image
                image = Image.open(io.BytesIO(data)).convert("RGB")
                text = pytesseract.image_to_string(image)
                if text.strip():
                    return text[:150_000], "OCR extracted", metadata
                return "", "No readable text", metadata
            except Exception:
                return "", "OCR unavailable (install Tesseract OCR)", metadata
    except Exception as exc:
        return "", f"Extraction failed: {type(exc).__name__}", metadata
    return "", "Unsupported file type", metadata


def extract_reference(text: str) -> str:
    patterns = [
        r"\b(?:reference\s*(?:no\.?|number)?|certificate\s*(?:no\.?|number)|registration\s*(?:no\.?|number)|ref)\s*[:#\-]?\s*([A-Z0-9][A-Z0-9/\-]{3,})",
        r"\b(DEMO-[A-Z0-9-]{4,})\b",
    ]
    for pattern in patterns:
        match = re.search(pattern, text or "", re.IGNORECASE)
        if match:
            return match.group(1).strip().rstrip(".,;")[:100]
    return ""


def parse_date(raw: str) -> date | None:
    raw = (raw or "").strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y", "%Y-%m-%d", "%d %B %Y", "%d %b %Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def find_expiry(text: str) -> tuple[date | None, str]:
    pattern = re.compile(
        r"(?:valid\s*(?:until|till|through|up\s*to)|expiry\s*date|expires?\s*(?:on|at)?|validity\s*(?:until|till|up\s*to))\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2}|\d{1,2}\s+[A-Za-z]+\s+\d{4})",
        re.IGNORECASE,
    )
    match = pattern.search(text or "")
    if not match:
        return None, ""
    return parse_date(match.group(1)), match.group(1)


def extract_institution_codes(text: str) -> list[str]:
    pattern = re.compile(
        r"(?:institution|institute|college)\s*(?:code|id|number|no\.?|identifier)|AICTE\s*(?:code|id)\s*[:#\-]?\s*([A-Z0-9][A-Z0-9/\-]{2,})",
        re.IGNORECASE,
    )
    results: list[str] = []
    # The explicit line parser is intentionally conservative to avoid flagging unrelated numbers.
    for line in (text or "").splitlines():
        if re.search(r"(?:institution|institute|college)\s*(?:code|id|number|no\.?|identifier)|AICTE\s*(?:code|id)", line, re.IGNORECASE):
            parts = re.split(r"[:#=\-]", line, maxsplit=1)
            if len(parts) == 2:
                value = re.search(r"[A-Z0-9][A-Z0-9/\-]{2,}", parts[1], re.IGNORECASE)
                if value:
                    results.append(value.group(0).strip())
    return results


def extract_institution_names(text: str) -> list[str]:
    results = []
    for line in (text or "").splitlines():
        match = re.search(r"(?:institution|institute|college)\s*name\s*[:=\-]\s*(.{4,120})", line, re.IGNORECASE)
        if match:
            results.append(match.group(1).strip())
    return results


def classify_document(text: str) -> tuple[str, int]:
    t = (text or "").casefold()
    patterns = {
        "Extension of Approval (EoA) letter": ["extension of approval", "eoa", "approved intake", "approval letter"],
        "Accreditation / recognition certificate": ["accreditation", "naac", "nba", "recognition certificate"],
        "Faculty roster": ["faculty roster", "faculty list", "staff name", "department", "faculty id"],
        "TER payment receipt": ["payment receipt", "transaction id", "ter payment", "fee receipt", "payment reference"],
        "Infrastructure evidence": ["infrastructure", "laboratory", "classroom", "building photo", "floor plan"],
        "Affidavit / undertaking": ["affidavit", "undertaking", "declaration by the institution"],
        "NOC / permission": ["no objection certificate", "noc", "permission letter"],
        "Fire safety certificate": ["fire safety", "fire noc", "fire department"],
        "Building stability certificate": ["building stability", "structural stability", "stability certificate"],
        "Faculty qualification evidence": ["degree certificate", "qualification certificate", "educational qualification"],
        "Land / building ownership or lease proof": ["lease deed", "ownership deed", "land ownership", "registered lease"],
    }
    scores = {kind: sum(1 for word in words if word in t) for kind, words in patterns.items()}
    best = max(scores, key=scores.get)
    return (best, scores[best]) if scores[best] else ("Unclassified", 0)


def parse_pdf_meta_date(value: str) -> datetime | None:
    if not value:
        return None
    match = re.search(r"D:(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?", value)
    if match:
        parts = [int(part or default) for part, default in zip(match.groups(), ["1", "1", "0", "0", "0", "0"])]
        try:
            return datetime(parts[0], parts[1], parts[2], parts[3], parts[4], parts[5])
        except ValueError:
            return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return None


def upsert_finding(db: sqlite3.Connection, app_id: str, finding: dict[str, Any]) -> None:
    now = utc_now()
    db.execute(
        """INSERT INTO findings(application_id, finding_key, doc_id, severity, title, detail, evidence,
           recommendation, status, is_active, created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'Open', 1, ?, ?)
           ON CONFLICT(application_id, finding_key) DO UPDATE SET
           doc_id=excluded.doc_id, severity=excluded.severity, title=excluded.title,
           detail=excluded.detail, evidence=excluded.evidence, recommendation=excluded.recommendation,
           status=CASE WHEN findings.status='Resolved' THEN 'Open' ELSE findings.status END,
           is_active=1, updated_at=excluded.updated_at""",
        (app_id, finding["key"], finding.get("doc_id"), finding["severity"], finding["title"],
         finding["detail"], finding.get("evidence", ""), finding.get("recommendation", ""), now, now),
    )


def run_validation(app_id: str, db: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_connection = db is None
    conn = db or sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    app_row = conn.execute("SELECT * FROM applications WHERE id=?", (app_id,)).fetchone()
    if not app_row:
        if own_connection:
            conn.close()
        raise ValueError("Application not found")
    application = dict(app_row)
    docs = [dict(row) for row in conn.execute("SELECT * FROM documents WHERE application_id=? ORDER BY id", (app_id,)).fetchall()]
    rules = get_rules()
    required = get_required_documents().get(application["category"], [])
    existing_types = {doc["doc_type"] for doc in docs}
    conn.execute("UPDATE findings SET is_active=0, updated_at=? WHERE application_id=?", (utc_now(), app_id))
    findings: list[dict[str, Any]] = []

    if rules.get("missing_documents", True):
        for doc_type in required:
            if doc_type not in existing_types:
                findings.append({
                    "key": f"missing:{doc_type}", "severity": "High", "title": f"Missing required document: {doc_type}",
                    "detail": f"The current {application['category']} checklist includes this item, but no file is attached under that document type.",
                    "evidence": f"Checklist rule: {application['category']} → {doc_type}",
                    "recommendation": "Upload the document or record an authorized exemption/alternative, where the applicable process permits it.",
                })

    seen_hashes: dict[str, dict[str, Any]] = {}
    for doc in docs:
        text = doc["extracted_text"] or ""
        doc_id = int(doc["id"])
        if rules.get("unreadable_documents", True) and len(text.strip()) < 20:
            findings.append({
                "key": f"unreadable:{doc_id}", "doc_id": doc_id, "severity": "Medium", "title": f"Text could not be reliably extracted: {doc['original_name']}",
                "detail": f"Extraction status: {doc['extraction_status']}. Automated checks may be incomplete for this file.",
                "evidence": "Extracted text length is below the review threshold.",
                "recommendation": "Review the file manually or upload a clearer, text-searchable copy. For scanned files, enable a supported OCR engine.",
            })

        if rules.get("file_signature", True) and not doc["signature_ok"]:
            findings.append({
                "key": f"signature:{doc_id}", "doc_id": doc_id, "severity": "High", "title": f"File signature does not match its extension: {doc['original_name']}",
                "detail": "The file content header does not match the declared file type. This can occur with a mislabeled or malformed file and is not, by itself, proof of manipulation.",
                "evidence": f"Declared extension: .{doc['extension'] or 'unknown'}; signature check failed.",
                "recommendation": "Re-upload the original file in its correct format and review the source copy.",
            })

        if rules.get("expired_documents", True):
            expiry, expiry_raw = find_expiry(text)
            if expiry and expiry < date.today():
                findings.append({
                    "key": f"expired:{doc_id}", "doc_id": doc_id, "severity": "High", "title": f"Possible expired document: {doc['original_name']}",
                    "detail": f"A validity phrase was followed by {expiry_raw}, which is before today's date ({date.today().strftime('%d %b %Y')}). Confirm the applicable validity rule and source record.",
                    "evidence": f"Extracted validity text: “valid until {expiry_raw}”.", 
                    "recommendation": "Verify the date against the issuing authority and request a current document if required.",
                })

        if rules.get("cross_document_consistency", True):
            for found_code in extract_institution_codes(text):
                if normalize(found_code) != normalize(application["institution_code"]):
                    findings.append({
                        "key": f"code-mismatch:{doc_id}:{normalize(found_code)}", "doc_id": doc_id, "severity": "High", "title": f"Institution identifier mismatch: {doc['original_name']}",
                        "detail": f"The application lists {application['institution_code']}, but the document contains {found_code}.",
                        "evidence": f"Application field: {application['institution_code']} · Extracted document field: {found_code}",
                        "recommendation": "Confirm the institution identifier and correct the application or provide the authoritative corrected document.",
                    })
            for found_name in extract_institution_names(text):
                ratio = SequenceMatcher(None, normalize(found_name), normalize(application["institution_name"])).ratio()
                if ratio < 0.68 and len(normalize(found_name)) > 6:
                    findings.append({
                        "key": f"name-mismatch:{doc_id}:{normalize(found_name)[:30]}", "doc_id": doc_id, "severity": "Medium", "title": f"Institution name differs from application: {doc['original_name']}",
                        "detail": f"The application name is “{application['institution_name']}”; the extracted document name is “{found_name}”. Differences may reflect abbreviations or OCR errors.",
                        "evidence": f"Name similarity indicator: {ratio:.0%} (heuristic, not a definitive identity check).",
                        "recommendation": "Confirm the legal institution name and inspect the source document before taking action.",
                    })

        if rules.get("duplicate_files", True):
            if doc["sha256"] in seen_hashes:
                first = seen_hashes[doc["sha256"]]
                findings.append({
                    "key": f"duplicate:{doc['sha256']}", "doc_id": doc_id, "severity": "Medium", "title": "Duplicate document detected",
                    "detail": f"This file has the same SHA-256 hash as “{first['original_name']}”. Duplicate uploads may be intentional, but should be checked.",
                    "evidence": f"SHA-256: {doc['sha256']}",
                    "recommendation": "Confirm whether the duplicate is expected and retain the correct version in the application record.",
                })
            else:
                seen_hashes[doc["sha256"]] = doc

        if rules.get("pdf_metadata_anomalies", True) and doc["extension"] == "pdf":
            meta = load_json_from_text(doc["metadata_json"])
            created = parse_pdf_meta_date(meta.get("creation_date", ""))
            modified = parse_pdf_meta_date(meta.get("modification_date", ""))
            if created and modified and modified < created:
                findings.append({
                    "key": f"metadata-time:{doc_id}", "doc_id": doc_id, "severity": "Low", "title": f"PDF metadata timestamps are inconsistent: {doc['original_name']}",
                    "detail": "The PDF metadata modification timestamp appears earlier than its creation timestamp. Metadata can be altered or copied during normal workflows; this is only a review indicator.",
                    "evidence": f"Creation metadata: {created.isoformat()} · Modification metadata: {modified.isoformat()}",
                    "recommendation": "Compare with the issuing authority's record or original digitally signed file.",
                })
            elif created and modified and (modified - created).days > 3650:
                findings.append({
                    "key": f"metadata-gap:{doc_id}", "doc_id": doc_id, "severity": "Low", "title": f"Large PDF metadata date gap: {doc['original_name']}",
                    "detail": "The creation and modification metadata differ by more than ten years. Metadata dates are not reliable proof of document authenticity.",
                    "evidence": f"Creation metadata: {created.date()} · Modification metadata: {modified.date()}",
                    "recommendation": "Review the source document and issuer record if the date history matters to this application.",
                })

        if rules.get("document_type_check", True):
            predicted_type, strength = classify_document(text)
            if strength >= 2 and predicted_type != doc["doc_type"]:
                findings.append({
                    "key": f"type-mismatch:{doc_id}:{normalize(predicted_type)}", "doc_id": doc_id, "severity": "Low", "title": f"Document category may not match its contents: {doc['original_name']}",
                    "detail": f"Selected type: {doc['doc_type']}. Keyword-based classification suggests: {predicted_type}. This is a lightweight heuristic and can be affected by document wording or OCR errors.",
                    "evidence": f"Classification keyword signals: {strength}.",
                    "recommendation": "Confirm the document type and reclassify the upload if needed.",
                })

        if doc["source_status"].startswith("Expired ("):
            findings.append({
                "key": f"issuer-expired:{doc_id}", "doc_id": doc_id, "severity": "High", "title": f"Issuer lookup returned an expired record: {doc['original_name']}",
                "detail": doc["source_detail"] or "The source response indicates that the record is expired.",
                "evidence": f"Source: {doc['source_name']}; status: {doc['source_status']}; reference: {doc['issuer_reference'] or 'not found'}",
                "recommendation": "Confirm the result against the actual issuing authority and request a current record if required.",
            })
        elif doc["source_status"].startswith("Mismatch ("):
            findings.append({
                "key": f"issuer-mismatch:{doc_id}", "doc_id": doc_id, "severity": "High", "title": f"Issuer lookup returned a record mismatch: {doc['original_name']}",
                "detail": doc["source_detail"] or "The synthetic source returned details that do not match the application.",
                "evidence": f"Source: {doc['source_name']}; status: {doc['source_status']}; reference: {doc['issuer_reference'] or 'not found'}",
                "recommendation": "Pause automated closure and route the discrepancy for authorized manual verification.",
            })

    for finding in findings:
        upsert_finding(conn, app_id, finding)

    active = conn.execute("SELECT severity, status FROM findings WHERE application_id=? AND is_active=1", (app_id,)).fetchall()
    open_findings = [row for row in active if row["status"] != "Resolved"]
    high_count = sum(1 for row in open_findings if row["severity"] in {"Critical", "High"})
    medium_count = sum(1 for row in open_findings if row["severity"] == "Medium")
    score = max(0, 100 - sum(SEVERITY_WEIGHT.get(row["severity"], 0) for row in open_findings))
    if not docs:
        status = "Draft — documents needed"
    elif high_count:
        status = "Needs official review"
    elif medium_count:
        status = "Review recommended"
    else:
        status = "Ready for official review"
    conn.execute("UPDATE applications SET status=?, readiness_score=?, updated_at=? WHERE id=?", (status, score, utc_now(), app_id))
    record_audit(conn, app_id, "Validation run", f"Rules evaluated: {len(rules)}; active findings: {len(open_findings)}; completeness & consistency score: {score}/100.")
    conn.commit()
    if own_connection:
        conn.close()
    return {"status": status, "score": score, "finding_count": len(open_findings)}


def load_json_from_text(value: str) -> dict[str, Any]:
    try:
        parsed = json.loads(value or "{}")
        return parsed if isinstance(parsed, dict) else {}
    except json.JSONDecodeError:
        return {}


def get_application(app_id: str) -> dict[str, Any]:
    row = get_db().execute("SELECT * FROM applications WHERE id=?", (app_id,)).fetchone()
    if not row:
        abort(404)
    return dict(row)


def get_application_documents(app_id: str) -> list[dict[str, Any]]:
    rows = get_db().execute("SELECT * FROM documents WHERE application_id=? ORDER BY id", (app_id,)).fetchall()
    docs = [dict(row) for row in rows]
    for doc in docs:
        doc["extracted_preview"] = (doc["extracted_text"] or "").strip()[:260]
        doc["metadata"] = load_json_from_text(doc.pop("metadata_json", "{}"))
        doc["size_label"] = f"{doc['size_bytes'] / 1024:.0f} KB" if doc["size_bytes"] >= 1024 else f"{doc['size_bytes']} B"
    return docs


def get_application_findings(app_id: str, active_only: bool = True) -> list[dict[str, Any]]:
    query = "SELECT * FROM findings WHERE application_id=?"
    if active_only:
        query += " AND is_active=1"
    query += " ORDER BY CASE severity WHEN 'Critical' THEN 0 WHEN 'High' THEN 1 WHEN 'Medium' THEN 2 WHEN 'Low' THEN 3 ELSE 4 END, created_at DESC"
    return [dict(row) for row in get_db().execute(query, (app_id,)).fetchall()]


def get_audit_logs(app_id: str, limit: int = 50) -> list[dict[str, Any]]:
    return [dict(row) for row in get_db().execute("SELECT * FROM audit_logs WHERE application_id=? ORDER BY id DESC LIMIT ?", (app_id, limit)).fetchall()]


def create_document_record(db: sqlite3.Connection, app_id: str, doc_type: str, filename: str, ext: str,
                           data: bytes, text: str, extraction_status: str, metadata: dict[str, Any],
                           stored_name: str | None, is_demo: bool = False, signature_ok: bool = True,
                           source_status: str = "Not checked", source_name: str = "", source_detail: str = "") -> int:
    reference = extract_reference(text)
    cur = db.execute(
        """INSERT INTO documents(application_id, doc_type, original_name, stored_name, extension, mime_type,
           size_bytes, sha256, extracted_text, extraction_status, metadata_json, source_status, source_name,
           source_detail, issuer_reference, signature_ok, is_demo, created_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (app_id, doc_type, filename[:255], stored_name, ext, mimetypes.guess_type(filename)[0] or "application/octet-stream",
         len(data), hashlib.sha256(data).hexdigest(), text[:150_000], extraction_status, json.dumps(metadata), source_status,
         source_name, source_detail, reference, int(signature_ok), int(is_demo), utc_now()),
    )
    return int(cur.lastrowid)


def issuer_lookup_for_document(db: sqlite3.Connection, app_id: str, doc_id: int) -> tuple[str, str, str]:
    doc_row = db.execute("SELECT * FROM documents WHERE id=? AND application_id=?", (doc_id, app_id)).fetchone()
    if not doc_row:
        raise ValueError("Document not found")
    doc = dict(doc_row)
    reference = doc["issuer_reference"] or extract_reference(doc["extracted_text"])
    if doc["is_demo"] and reference in SYNTHETIC_ISSUER_REGISTRY:
        record = SYNTHETIC_ISSUER_REGISTRY[reference]
        status = record["status"]
        detail = f"{record['note']} Returned institution code: {record['institution_code']}; valid until: {record['valid_until']}. SIMULATED TEST DATA — not an official source response."
        source = record["issuer"]
        if normalize(record["institution_code"]) != normalize(get_application(app_id)["institution_code"]):
            status = "Mismatch (synthetic source)"
            detail += " The returned institution code differs from the application record."
        db.execute("UPDATE documents SET source_status=?, source_name=?, source_detail=?, issuer_reference=? WHERE id=?",
                   (status, source, detail, reference, doc_id))
        record_audit(db, app_id, "Synthetic issuer lookup", f"Document {doc_id}; reference {reference}; result {status}. No live authority was contacted.")
        db.commit()
        return status, source, detail
    status = "Manual verification required"
    source = "No authorized source connected"
    detail = "No live issuer API or authorized portal is configured in this prototype. The reference has been captured for manual confirmation; no verification is claimed."
    db.execute("UPDATE documents SET source_status=?, source_name=?, source_detail=?, issuer_reference=? WHERE id=?",
               (status, source, detail, reference, doc_id))
    record_audit(db, app_id, "Manual verification routed", f"Document {doc_id}; reference {reference or 'not detected'}; no authorized live source configured.")
    db.commit()
    return status, source, detail


@app.route("/")
def dashboard():
    db = get_db()
    applications = [dict(row) for row in db.execute("SELECT * FROM applications ORDER BY updated_at DESC LIMIT 8").fetchall()]
    stats = {
        "applications": db.execute("SELECT COUNT(*) FROM applications").fetchone()[0],
        "needs_review": db.execute("SELECT COUNT(*) FROM applications WHERE status='Needs official review'").fetchone()[0],
        "pending_source": db.execute("SELECT COUNT(*) FROM documents WHERE source_status IN ('Not checked','Manual verification required')").fetchone()[0],
        "open_findings": db.execute("SELECT COUNT(*) FROM findings WHERE is_active=1 AND status!='Resolved'").fetchone()[0],
    }
    queue = [dict(row) for row in db.execute(
        """SELECT f.*, a.institution_name, a.institution_code FROM findings f JOIN applications a ON a.id=f.application_id
           WHERE f.is_active=1 AND f.status!='Resolved' ORDER BY CASE f.severity WHEN 'Critical' THEN 0 WHEN 'High' THEN 1 WHEN 'Medium' THEN 2 ELSE 3 END, f.id DESC LIMIT 6"""
    ).fetchall()]
    return render_template("dashboard.html", applications=applications, stats=stats, queue=queue)


@app.route("/applications/new")
def new_application():
    return render_template("new_application.html", required_docs=get_required_documents(), allowed_extensions=", ".join(sorted(ALLOWED_EXTENSIONS)))


@app.route("/applications", methods=["POST"])
def create_application():
    institution_name = request.form.get("institution_name", "").strip()[:180]
    institution_code = request.form.get("institution_code", "").strip()[:80]
    category = request.form.get("category", "").strip()
    academic_year = request.form.get("academic_year", "").strip()[:20]
    contact_email = request.form.get("contact_email", "").strip()[:180]
    categories = set(get_required_documents())
    if not institution_name or not institution_code or not category or not academic_year:
        flash("Complete the institution name, identifier, application category and academic year.", "error")
        return redirect(url_for("new_application"))
    if category not in categories:
        flash("Choose a supported application category.", "error")
        return redirect(url_for("new_application"))
    files = request.files.getlist("documents")
    doc_types = request.form.getlist("document_types")
    selected = [(f, doc_types[i] if i < len(doc_types) else "Other supporting document") for i, f in enumerate(files) if f and f.filename]
    if len(selected) > app.config["MAX_FILES_PER_APPLICATION"]:
        flash(f"Upload up to {app.config['MAX_FILES_PER_APPLICATION']} documents at a time.", "error")
        return redirect(url_for("new_application"))
    for file, _doc_type in selected:
        if not allowed_file(file.filename):
            flash(f"Unsupported file type: {secure_filename(file.filename)}. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}.", "error")
            return redirect(url_for("new_application"))

    app_id = f"AG-{date.today().strftime('%Y%m%d')}-{uuid.uuid4().hex[:5].upper()}"
    db = get_db()
    now = utc_now()
    db.execute(
        "INSERT INTO applications(id, institution_name, institution_code, category, academic_year, contact_email, status, readiness_score, created_at, updated_at, is_demo) VALUES (?, ?, ?, ?, ?, ?, 'Processing', 100, ?, ?, 0)",
        (app_id, institution_name, institution_code, category, academic_year, contact_email, now, now),
    )
    record_audit(db, app_id, "Application created", f"Category: {category}; academic year: {academic_year}.")
    for file, doc_type in selected:
        original = secure_filename(file.filename) or f"document.{file.filename.rsplit('.', 1)[-1].lower()}"
        ext = original.rsplit(".", 1)[-1].lower()
        data = file.read()
        if not data:
            continue
        stored = f"{uuid.uuid4().hex}_{original}"
        target = UPLOAD_DIR / stored
        target.write_bytes(data)
        text, extraction_status, metadata = extract_document(data, ext)
        signature_ok = file_signature_ok(ext, data)
        create_document_record(db, app_id, doc_type or "Other supporting document", original, ext, data, text,
                               extraction_status, metadata, stored, False, signature_ok)
        record_audit(db, app_id, "Document uploaded", f"{original}; type: {doc_type}; extraction: {extraction_status}; {len(data)} bytes.")
    db.commit()
    result = run_validation(app_id, db)
    flash(f"Application created. Preliminary checks found {result['finding_count']} item(s) for review.", "success")
    return redirect(url_for("application_detail", app_id=app_id))


@app.route("/applications/<app_id>")
def application_detail(app_id: str):
    application = get_application(app_id)
    docs = get_application_documents(app_id)
    findings = get_application_findings(app_id)
    logs = get_audit_logs(app_id)
    required = get_required_documents().get(application["category"], [])
    present = {doc["doc_type"] for doc in docs}
    checklist = [{"name": name, "present": name in present} for name in required]
    summary = {
        "documents": len(docs),
        "required": len(required),
        "present_required": sum(1 for item in checklist if item["present"]),
        "open": sum(1 for f in findings if f["status"] != "Resolved"),
        "high": sum(1 for f in findings if f["status"] != "Resolved" and f["severity"] in {"Critical", "High"}),
        "verified": sum(1 for doc in docs if doc["source_status"].startswith("Verified (synthetic") or doc["source_status"] == "Verified (authorized source)"),
        "pending": sum(1 for doc in docs if doc["source_status"] in {"Not checked", "Manual verification required"}),
    }
    return render_template("application_detail.html", application=application, docs=docs, findings=findings,
                           logs=logs, checklist=checklist, summary=summary)


@app.route("/applications/<app_id>/validate", methods=["POST"])
def validate_application(app_id: str):
    get_application(app_id)
    result = run_validation(app_id, get_db())
    flash(f"Validation completed: {result['finding_count']} active item(s); score {result['score']}/100.", "success")
    return redirect(url_for("application_detail", app_id=app_id) + "#findings")


@app.route("/applications/<app_id>/documents/<int:doc_id>/issuer-check", methods=["POST"])
def issuer_check(app_id: str, doc_id: int):
    get_application(app_id)
    try:
        status, _source, _detail = issuer_lookup_for_document(get_db(), app_id, doc_id)
        result = run_validation(app_id, get_db())
        flash(f"Source workflow status: {status}. {result['finding_count']} active item(s) remain.", "success" if status.startswith("Verified") else "warning")
    except ValueError:
        abort(404)
    return redirect(url_for("application_detail", app_id=app_id) + "#documents")


@app.route("/applications/<app_id>/findings/<int:finding_id>/review", methods=["POST"])
def review_finding(app_id: str, finding_id: int):
    get_application(app_id)
    action = request.form.get("review_action", "").strip()
    note = request.form.get("review_note", "").strip()[:1000]
    status_map = {"acknowledge": "Acknowledged", "resolve": "Resolved", "reopen": "Open"}
    if action not in status_map:
        flash("Choose a valid review action.", "error")
        return redirect(url_for("application_detail", app_id=app_id) + "#findings")
    if action == "resolve" and len(note) < 5:
        flash("Add a short review note before resolving a finding.", "error")
        return redirect(url_for("application_detail", app_id=app_id) + "#findings")
    db = get_db()
    cur = db.execute("UPDATE findings SET status=?, updated_at=? WHERE id=? AND application_id=? AND is_active=1",
                     (status_map[action], utc_now(), finding_id, app_id))
    if cur.rowcount == 0:
        abort(404)
    record_audit(db, app_id, f"Finding {status_map[action].lower()}", f"Finding #{finding_id}. {note}")
    db.commit()
    # Recompute the score/status without resetting reviewer states.
    rows = db.execute("SELECT severity, status FROM findings WHERE application_id=? AND is_active=1", (app_id,)).fetchall()
    open_rows = [row for row in rows if row["status"] != "Resolved"]
    score = max(0, 100 - sum(SEVERITY_WEIGHT.get(row["severity"], 0) for row in open_rows))
    if not get_application_documents(app_id):
        app_status = "Draft — documents needed"
    elif any(row["severity"] in {"Critical", "High"} for row in open_rows):
        app_status = "Needs official review"
    elif open_rows:
        app_status = "Review recommended"
    else:
        app_status = "Ready for official review"
    db.execute("UPDATE applications SET status=?, readiness_score=?, updated_at=? WHERE id=?", (app_status, score, utc_now(), app_id))
    db.commit()
    flash(f"Finding marked {status_map[action].lower()}.", "success")
    return redirect(url_for("application_detail", app_id=app_id) + "#findings")


@app.route("/applications/<app_id>/report")
def report_view(app_id: str):
    application = get_application(app_id)
    docs = get_application_documents(app_id)
    findings = [item for item in get_application_findings(app_id) if item["status"] != "Resolved"]
    logs = get_audit_logs(app_id, 100)
    required = get_required_documents().get(application["category"], [])
    present = {doc["doc_type"] for doc in docs}
    checklist = [{"name": name, "present": name in present} for name in required]
    return render_template("report.html", application=application, docs=docs, findings=findings, logs=logs, checklist=checklist, generated_at=utc_now())


@app.route("/applications/<app_id>/report.json")
def report_json(app_id: str):
    application = get_application(app_id)
    docs = get_application_documents(app_id)
    findings = get_application_findings(app_id)
    payload = {
        "product": "AegisAI prototype",
        "application": {key: application[key] for key in ["id", "institution_name", "institution_code", "category", "academic_year", "status", "readiness_score", "created_at", "updated_at"]},
        "documents": [{key: doc.get(key) for key in ["id", "doc_type", "original_name", "size_bytes", "sha256", "extraction_status", "source_status", "source_name", "source_detail", "issuer_reference"]} for doc in docs],
        "findings": [{key: f.get(key) for key in ["id", "severity", "title", "detail", "evidence", "recommendation", "status"]} for f in findings],
        "audit_log": [{key: log.get(key) for key in ["actor", "action", "detail", "created_at"]} for log in get_audit_logs(app_id, 100)],
        "disclaimer": "Decision-support prototype only. No live AICTE or issuer system was contacted unless explicitly stated; synthetic source results are labelled.",
    }
    record_audit(get_db(), app_id, "JSON report exported", "Evidence-linked application report exported.")
    get_db().commit()
    response = make_response(json.dumps(payload, indent=2, ensure_ascii=False))
    response.headers["Content-Type"] = "application/json; charset=utf-8"
    response.headers["Content-Disposition"] = f"attachment; filename={app_id}-aegisai-report.json"
    return response


@app.route("/applications/<app_id>/report.csv")
def report_csv(app_id: str):
    application = get_application(app_id)
    findings = get_application_findings(app_id)
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["Application ID", "Institution", "Institution Code", "Category", "Academic Year", "Application Status", "Readiness Score", "Finding ID", "Severity", "Finding", "Evidence", "Recommendation", "Review Status"])
    if findings:
        for finding in findings:
            writer.writerow([application["id"], application["institution_name"], application["institution_code"], application["category"], application["academic_year"], application["status"], application["readiness_score"], finding["id"], finding["severity"], finding["title"], finding["evidence"], finding["recommendation"], finding["status"]])
    else:
        writer.writerow([application["id"], application["institution_name"], application["institution_code"], application["category"], application["academic_year"], application["status"], application["readiness_score"], "", "", "No active findings", "", "", ""])
    record_audit(get_db(), app_id, "CSV report exported", "Evidence-linked finding report exported.")
    get_db().commit()
    response = make_response("\ufeff" + output.getvalue())
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    response.headers["Content-Disposition"] = f"attachment; filename={app_id}-aegisai-findings.csv"
    return response


@app.route("/rules", methods=["GET", "POST"])
def rules_page():
    if request.method == "POST":
        current = get_rules()
        updated = {key: request.form.get(key) == "on" for key in current}
        RULES_PATH.write_text(json.dumps(updated, indent=2), encoding="utf-8")
        record_audit(get_db(), None, "Validation rules updated", ", ".join(f"{key}={'on' if value else 'off'}" for key, value in updated.items()))
        get_db().commit()
        flash("Rule settings saved for this local prototype.", "success")
        return redirect(url_for("rules_page"))
    return render_template("rules.html", rules=get_rules(), required_docs=get_required_documents())


@app.route("/demo", methods=["POST"])
def load_demo():
    db = get_db()
    app_id = f"DEMO-{date.today().strftime('%Y%m%d')}-{uuid.uuid4().hex[:4].upper()}"
    now = utc_now()
    institution = "Northstar Institute of Technology"
    code = "AICTE-DEMO-042"
    category = "Extension of Approval (EoA)"
    db.execute(
        "INSERT INTO applications(id, institution_name, institution_code, category, academic_year, contact_email, status, readiness_score, created_at, updated_at, is_demo) VALUES (?, ?, ?, ?, ?, ?, 'Processing', 100, ?, ?, 1)",
        (app_id, institution, code, category, "2026-27", "demo@example.invalid", now, now),
    )
    demo_docs = [
        ("Extension of Approval (EoA) letter", "demo_eoa_letter.txt", "NORTHSTAR INSTITUTE OF TECHNOLOGY\nInstitution Code: AICTE-DEMO-042\nDocument: Extension of Approval (EoA)\nAcademic Year: 2026-27\nReference No: DEMO-EOA-001\nIssued on: 15/05/2026\nValid until: 31/03/2027\nApproved programme: Computer Science\nThis is a synthetic sample document for prototype demonstration only."),
        ("Accreditation / recognition certificate", "demo_accreditation.txt", "NORTHSTAR INSTITUTE OF TECHNOLOGY\nInstitution Code: AICTE-DEMO-042\nDocument: Accreditation / Recognition Certificate\nReference No: DEMO-ACC-EXPIRED\nIssued on: 01/04/2022\nValid until: 31/03/2025\nThis is a synthetic sample document for prototype demonstration only."),
        ("Faculty roster", "demo_faculty_roster.txt", "NORTHSTAR INSTITUTE OF TECHNOLOGY\nInstitution Code: AICTE-DEMO-999\nDocument: Faculty Roster\nAcademic Year: 2026-27\nFaculty roster: 34 faculty members listed.\nThis is a synthetic sample document for prototype demonstration only."),
        ("TER payment receipt", "demo_ter_receipt.txt", "NORTHSTAR INSTITUTE OF TECHNOLOGY\nInstitution Code: AICTE-DEMO-042\nDocument: TER Payment Receipt\nReference No: DEMO-TER-001\nTransaction ID: SYNTH-TER-7781\nPaid on: 10/06/2026\nThis is a synthetic sample document for prototype demonstration only."),
        ("Infrastructure evidence", "demo_infrastructure.txt", "NORTHSTAR INSTITUTE OF TECHNOLOGY\nInstitution Code: AICTE-DEMO-042\nDocument: Infrastructure Evidence\nLaboratory: 4; Classrooms: 18; Library and workshop evidence attached.\nThis is a synthetic text placeholder for an infrastructure evidence file."),
    ]
    for doc_type, filename, content in demo_docs:
        data = content.encode("utf-8")
        ref = extract_reference(content)
        if ref in SYNTHETIC_ISSUER_REGISTRY:
            registry = SYNTHETIC_ISSUER_REGISTRY[ref]
            status = registry["status"]
            source = registry["issuer"]
            detail = f"{registry['note']} Returned institution code: {registry['institution_code']}; valid until: {registry['valid_until']}. SIMULATED TEST DATA — not an official source response."
        else:
            status, source, detail = "Not checked", "", ""
        create_document_record(db, app_id, doc_type, filename, "txt", data, content, "Text extracted", {}, None,
                               True, True, status, source, detail)
    record_audit(db, app_id, "Synthetic demo case created", "All records are fabricated for demonstration. Synthetic issuer lookup results are not live AICTE or issuer verification. One required affidavit/undertaking is intentionally absent.")
    db.commit()
    result = run_validation(app_id, db)
    flash(f"Synthetic demo loaded with {result['finding_count']} review item(s). No real institutional data or live source was used.", "success")
    return redirect(url_for("application_detail", app_id=app_id))


@app.route("/healthz")
def healthz():
    return jsonify({"status": "ok", "service": "AegisAI SIH26261 prototype", "database": "local SQLite"})


@app.errorhandler(413)
def too_large(_error):
    return render_template("error.html", code=413, message="The upload is larger than the 20 MB request limit. Split the documents and try again."), 413


@app.errorhandler(400)
def bad_request(error):
    return render_template("error.html", code=400, message=getattr(error, "description", "The request could not be processed.")), 400


@app.errorhandler(404)
def not_found(_error):
    return render_template("error.html", code=404, message="The page or application you requested was not found."), 404


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), debug=False)
